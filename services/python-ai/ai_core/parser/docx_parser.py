from docx import Document


class DOCXParser:

    def parse(self, file_path: str, *, metadata: dict | None = None):

        metadata = dict(metadata or {})

        document = Document(file_path)

        blocks = []

        # --------------------------
        # Параграфы
        # --------------------------

        for index, paragraph in enumerate(
            document.paragraphs
        ):

            text = paragraph.text.strip()

            if not text:
                continue

            blocks.append({

                "type": "paragraph",

                "index": index,

                "text": text,
                "source_kind": "paragraph",
                "page": None,
                "bbox": None,
            })

        # --------------------------
        # Таблицы
        # --------------------------

        tables = []

        for table_index, table in enumerate(
            document.tables
        ):

            rows = []

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                rows.append(cells)

            tables.append({

                "table": table_index,

                "rows": rows,
                "source_kind": "table",
            })

        full_text = "\n".join(
            block["text"]
            for block in blocks
        )

        return {

            "type": "docx",

            "schema_version": metadata.get("schema_version", "parser-v2"),

            "source": metadata.get("source") or file_path,

            "stage": metadata.get("stage"),

            "section": metadata.get("section"),

            "text": full_text,

            "blocks": blocks,

            "tables": tables
        }
