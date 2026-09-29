import sys
import unittest
from io import BytesIO
import json
from pathlib import Path


SERVICE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SERVICE_ROOT / "ai_core"))
sys.path.insert(0, str(SERVICE_ROOT))

from chunker import DocumentChunker  # noqa: E402
from search import ParameterExtractor, SourceAwareReranker  # noqa: E402
from app import process_document  # noqa: E402


class PipelineRegressionTests(unittest.TestCase):
    def test_chunker_keeps_drawing_dimensions_separate(self):
        document = {
            "type": "pdf",
            "schema_version": "parser-v2",
            "source": "plan.pdf",
            "stage": "RD",
            "pages": [{
                "page": 12,
                "blocks": [
                    {"bbox": [10, 10, 100, 20], "text": "500 1700 425 500", "block_type": 0, "source_kind": "text"},
                    {"bbox": [10, 80, 500, 100], "text": "Строительный объем — 42 560 м³", "block_type": 0, "source_kind": "text"},
                ],
            }],
        }
        chunks = DocumentChunker(max_words=48).chunk(document)
        self.assertEqual(len(chunks["chunks"]), 2)
        self.assertEqual(chunks["chunks"][0]["content_kind"], "drawing_dimension")
        self.assertFalse(chunks["chunks"][0]["searchable"])
        self.assertEqual(chunks["chunks"][1]["page"], 12)
        self.assertEqual(chunks["chunks"][1]["source"], "plan.pdf")
        self.assertEqual(chunks["chunks"][1]["stage"], "RD")
        self.assertEqual(chunks["chunks"][1]["bbox"], [10.0, 80.0, 500.0, 100.0])

    def test_reranker_rejects_number_soup_and_accepts_parameter_anchor(self):
        reranker = SourceAwareReranker()
        result = reranker.rerank([
            {
                "id": "drawing",
                "text": "500 1700 425 500",
                "content_kind": "drawing_dimension",
                "searchable": False,
                "semantic_score": 0.99,
            },
            {
                "id": "tep",
                "text": "Строительный объем здания — 42 560 м³",
                "content_kind": "text",
                "searchable": True,
                "semantic_score": 0.72,
                "page": 12,
                "bbox": [10, 80, 500, 100],
            },
        ], {
            "code": "M-004",
            "parameter": "Строительный объем (Общий)",
            "unit": "м³",
            "section": "Раздел 1. ПЗ",
            "source": "Раздел ПЗ: Таблица ТЭП",
            "stage": "PD",
        })
        self.assertEqual([item["id"] for item in result["results"]], ["tep"])
        self.assertEqual(result["rejected"][0]["reason"], "non_text_layout_block")

    def test_extractor_extracts_m004_value(self):
        extractor = ParameterExtractor()
        result = extractor.extract(
            {"code": "M-004", "parameter": "Строительный объем (Общий)", "unit": "м³"},
            [{"id": "tep", "text": "Строительный объем здания — 42 560 м³", "rerank_score": 0.9}],
        )
        self.assertEqual(result[0]["extracted_value"], "42 560 м³")
        self.assertEqual(result[0]["normalized_value"], "42560|м3")
        self.assertEqual(result[0]["extraction_status"], "CONFIRMED")

    def test_extractor_extracts_concrete_class(self):
        extractor = ParameterExtractor()
        result = extractor.extract(
            {
                "code": "M-061",
                "parameter": "Класс прочности бетона монолитных конструкций",
                "unit": "Марка (B)",
            },
            [{"id": "spec", "text": "Бетон монолитных конструкций класса B25", "rerank_score": 0.9}],
        )
        self.assertEqual(result[0]["normalized_value"], "B25")

    def test_extractor_returns_no_confirmed_value(self):
        extractor = ParameterExtractor()
        result = extractor.extract(
            {"code": "M-004", "parameter": "Строительный объем (Общий)", "unit": "м³"},
            [{"id": "text", "text": "Устройство бетонной подготовки и гидроизоляции", "rerank_score": 0.6}],
        )
        self.assertEqual(result, [])

    def test_process_document_returns_confirmed_values_and_not_found(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Строительный объем здания — 42 560 м³")
        document.add_paragraph("Бетон монолитных конструкций класса B25")
        stream = BytesIO()
        document.save(stream)

        queries = [
            {
                "code": "M-004",
                "parameter": "Строительный объем (Общий)",
                "unit": "м³",
                "section": "Раздел 1. ПЗ",
                "source": "Раздел ПЗ: Таблица ТЭП",
            },
            {
                "code": "M-061",
                "parameter": "Класс прочности бетона монолитных конструкций",
                "unit": "Марка (B)",
                "section": "Раздел 5. КР",
                "source": "Общие данные; Спецификация материалов (КР)",
            },
            {
                "code": "M-999",
                "parameter": "Площадь застройки",
                "unit": "м²",
                "section": "Раздел 1. ПЗ",
                "source": "Раздел ПЗУ: Таблица ТЭП",
            },
        ]
        result = process_document(
            {
                "object_id": "obj-test",
                "process_id": "process-test",
                "stage": "PD",
                "queries": json.dumps(queries, ensure_ascii=False),
            },
            {
                "name": "synthetic.docx",
                "content": stream.getvalue(),
                "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        self.assertEqual(result["model"]["semantic_top_k"], 20)
        m004 = result["results_by_parameter"]["M-004"][0]
        self.assertEqual(m004["normalized_value"], "42560|м3")
        self.assertEqual(m004["source"], "synthetic.docx")
        self.assertEqual(m004["stage"], "PD")
        self.assertEqual(m004["section"], "Раздел 1. ПЗ")
        self.assertEqual(result["results_by_parameter"]["M-061"][0]["normalized_value"], "B25")
        self.assertEqual(result["parameter_status"]["M-999"]["status"], "NOT_FOUND")


if __name__ == "__main__":
    unittest.main()
